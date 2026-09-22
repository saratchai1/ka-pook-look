from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.dml.color import RGBColor
import os
import argparse

parser=argparse.ArgumentParser(description='Restyle the Present Asset PowerPoint without changing text/table content.')
parser.add_argument('src', help='Source .pptx')
parser.add_argument('-o','--out', default='Present Asset - Redesigned.pptx', help='Output .pptx')
args=parser.parse_args()
SRC = args.src
OUT = args.out

NAVY = RGBColor(10, 43, 75)
BLUE = RGBColor(31, 105, 180)
CYAN = RGBColor(30, 162, 209)
TEAL = RGBColor(0, 139, 162)
INK = RGBColor(25, 42, 58)
LIGHT2 = RGBColor(226, 239, 249)
LIGHT3 = RGBColor(248, 251, 254)
BORDER = RGBColor(202, 221, 236)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(238, 242, 246)
PALE_BLUE = RGBColor(232, 244, 253)
PALE_CYAN = RGBColor(231, 248, 250)
PALE_NAVY = RGBColor(218, 231, 242)
FONT = 'Arial'

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height

def inch(v): return Inches(v)

def snapshot_original(prs):
    snap={}
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False):
                snap[(si, sh.shape_id, "text")]=sh.text
            if getattr(sh, "has_table", False):
                for r,row in enumerate(sh.table.rows):
                    for c,cell in enumerate(row.cells):
                        snap[(si, sh.shape_id, f"cell:{r}:{c}")]=cell.text
    return snap

BEFORE = snapshot_original(prs)
ORIG = [list(slide.shapes) for slide in prs.slides]

def style_title(shape, size=24):
    shape.left=inch(0.78); shape.top=inch(0.80); shape.width=inch(11.85); shape.height=inch(0.62)
    if not getattr(shape,'has_text_frame',False): return
    tf=shape.text_frame
    tf.word_wrap=True
    tf.margin_left=tf.margin_right=Pt(0)
    tf.margin_top=tf.margin_bottom=Pt(0)
    tf.vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        p.alignment=PP_ALIGN.LEFT
        p.space_before=Pt(0); p.space_after=Pt(0)
        try: p.line_spacing=1.0
        except: pass
        for r in p.runs:
            r.font.name=FONT; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=NAVY

def add_rect(slide, x,y,w,h, fill, line=None, radius=False):
    shp=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                               inch(x),inch(y),inch(w),inch(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb=line; shp.line.width=Pt(0.8)
    return shp

def add_panel_and_accent(slide, with_title=True):
    panel=add_rect(slide,0.28,0.70,12.77,6.57,WHITE,BORDER,True)
    add_rect(slide,0.46,0.70,1.55,0.055,BLUE,None,False)
    add_rect(slide,2.04,0.70,0.42,0.055,CYAN,None,False)
    if with_title:
        add_rect(slide,0.58,0.83,0.07,0.47,BLUE,None,False)
    return panel

def add_card(slide,x,y,w,h,fill=LIGHT3,line=BORDER):
    return add_rect(slide,x,y,w,h,fill,line,True)

def move_original_content_to_front(slide, original_shapes):
    for sh in original_shapes:
        decorative = False
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            decorative = True
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and (not getattr(sh,'has_text_frame',False) or not sh.text.strip()):
            decorative = True
        if not decorative:
            el=sh._element
            try:
                slide.shapes._spTree.remove(el)
                slide.shapes._spTree.append(el)
            except Exception:
                pass

def resize_table(shape,x,y,w,h):
    t=shape.table
    old_c=[c.width for c in t.columns]
    old_r=[r.height for r in t.rows]
    sumc=sum(old_c) or 1; sumr=sum(old_r) or 1
    shape.left=inch(x); shape.top=inch(y); shape.width=inch(w); shape.height=inch(h)
    totalw=inch(w); totalh=inch(h)
    for i,c in enumerate(t.columns):
        c.width=int(totalw*old_c[i]/sumc)
    avg=totalh/len(t.rows)
    if max(old_r)/(min([v for v in old_r if v>0]) or 1) > 2.5:
        for r in t.rows: r.height=int(avg)
    else:
        for i,r in enumerate(t.rows): r.height=int(totalh*old_r[i]/sumr)

def style_table2(shape, font_size=9, header_rows=1, policy=False, proposed_cols=None, old_cols=None):
    t=shape.table; rows=len(t.rows); cols=len(t.columns)
    proposed_cols=proposed_cols or []; old_cols=old_cols or []
    for r in range(rows):
        row0=t.cell(r,0).text.strip() if cols else ''
        rowtxt=' '.join(t.cell(r,c).text.strip() for c in range(cols))
        is_final=(r==rows-1 and any(k in rowtxt for k in ['ยอดรวม','รวมทั้งหมด']))
        is_subtotal=(not is_final and row0.startswith('รวม '))
        for c in range(cols):
            cell=t.cell(r,c)
            cell.margin_left=inch(0.025); cell.margin_right=inch(0.025); cell.margin_top=inch(0.012); cell.margin_bottom=inch(0.012)
            cell.vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE
            if r<header_rows:
                fill=NAVY if r==0 else LIGHT2; fontc=WHITE if r==0 else NAVY; bold=True
            elif is_final:
                fill=NAVY; fontc=WHITE; bold=True
            elif is_subtotal:
                fill=PALE_NAVY; fontc=NAVY; bold=True
            else:
                fill=WHITE if r%2 else LIGHT3; fontc=INK; bold=False
                if policy and c in old_cols: fill=GRAY
                if policy and c in proposed_cols: fill=PALE_BLUE
            cell.fill.solid(); cell.fill.fore_color.rgb=fill
            for p in cell.text_frame.paragraphs:
                p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.CENTER
                p.space_before=Pt(0); p.space_after=Pt(0)
                try: p.line_spacing=0.9
                except: pass
                for run in p.runs:
                    run.font.name=FONT; run.font.size=Pt(font_size); run.font.color.rgb=fontc; run.font.bold=bold

def set_pic(shape,x,y,w,h):
    shape.left=inch(x); shape.top=inch(y); shape.width=inch(w); shape.height=inch(h)

def style_body(shape, size=13, color=INK, bold_first=False):
    if not getattr(shape,'has_text_frame',False): return
    tf=shape.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(3); tf.margin_bottom=Pt(3)
    for i,p in enumerate(tf.paragraphs):
        p.space_before=Pt(0); p.space_after=Pt(4)
        try: p.line_spacing=1.05
        except: pass
        for r in p.runs:
            r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=color
            if bold_first and i==0: r.font.bold=True

for i,slide in enumerate(prs.slides,1):
    orig=ORIG[i-1]
    add_panel_and_accent(slide, with_title=(i not in [1,19]))
    move_original_content_to_front(slide,orig)

s=prs.slides[0]
add_card(s,0.78,1.35,11.78,5.45,WHITE,BORDER)
for sh in ORIG[0]:
    if sh.name=='TextBox 15':
        sh.left=inch(1.05); sh.top=inch(1.55); sh.width=inch(11.05); sh.height=inch(4.95)
        tf=sh.text_frame; tf.word_wrap=True
        tf.margin_left=Pt(2); tf.margin_right=Pt(2); tf.margin_top=Pt(2); tf.margin_bottom=Pt(2)
        for p in tf.paragraphs:
            p.space_before=Pt(0); p.space_after=Pt(7)
            try: p.line_spacing=1.02
            except: pass
            for r in p.runs:
                r.font.name=FONT; r.font.size=Pt(20); r.font.color.rgb=NAVY; r.font.bold=False
add_rect(s,0.88,1.56,0.055,4.58,BLUE,None)
for sh in [x for x in ORIG[0] if x.name=='TextBox 15']:
    el=sh._element; s.shapes._spTree.remove(el); s.shapes._spTree.append(el)

s=prs.slides[1]; o=ORIG[1]
style_title(o[0],23)
body=o[1]; body.left=inch(0.82); body.top=inch(1.38); body.width=inch(11.85); body.height=inch(0.85); style_body(body,12.2)
resize_table(o[2],0.48,2.30,12.37,4.67); style_table2(o[2],7.6,header_rows=2)
for c in range(12):
    cell=o[2].table.cell(0,c)
    fill = NAVY if c<4 else BLUE if c<8 else TEAL
    cell.fill.solid(); cell.fill.fore_color.rgb=fill
    for p in cell.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb=WHITE; r.font.bold=True

s=prs.slides[2]; o=ORIG[2]
style_title(o[0],23)
resize_table(o[3],0.55,1.48,6.35,5.38); style_table2(o[3],7.7,1)
add_card(s,7.12,1.45,5.63,2.80,WHITE,BORDER); set_pic(o[4],7.32,1.64,5.23,2.42)
lbl=o[1]; lbl.left=inch(7.18); lbl.top=inch(4.38); lbl.width=inch(5.52); lbl.height=inch(0.36); style_body(lbl,11.5,NAVY)
for p in lbl.text_frame.paragraphs:
    for r in p.runs: r.font.bold=True
resize_table(o[2],7.18,4.78,5.52,1.97); style_table2(o[2],7.1,1)

s=prs.slides[3]; o=ORIG[3]
style_title(o[0],24)
resize_table(o[1],0.72,1.52,5.40,5.35); style_table2(o[1],8.4,1)
add_card(s,6.35,1.52,6.35,4.22,WHITE,BORDER); set_pic(o[2],6.58,1.78,5.90,3.64)

s=prs.slides[4]; o=ORIG[4]
style_title(o[1],24)
resize_table(o[4],0.72,1.52,5.40,5.30); style_table2(o[4],8.5,1)
add_card(s,6.35,1.52,6.35,4.35,WHITE,BORDER); set_pic(o[5],6.56,1.77,5.93,3.72)

s=prs.slides[5]; o=ORIG[5]
style_title(o[1],21)
resize_table(o[4],1.15,1.58,11.00,4.45); style_table2(o[4],10.0,1)
add_card(s,1.15,6.15,11.00,0.66,PALE_BLUE,BORDER)
note=o[5]; note.left=inch(1.30); note.top=inch(6.28); note.width=inch(10.70); note.height=inch(0.38); style_body(note,10.8,NAVY)
for p in note.text_frame.paragraphs:
    for r in p.runs: r.font.bold=True

s=prs.slides[6]; o=ORIG[6]
style_title(o[1],23)
resize_table(o[0],0.48,1.47,12.36,5.62); style_table2(o[0],6.8,1)

s=prs.slides[7]; o=ORIG[7]
style_title(o[0],24)
add_card(s,0.78,1.50,11.78,5.35,LIGHT3,BORDER)
body=o[1]; body.left=inch(0.98); body.top=inch(1.70); body.width=inch(11.38); body.height=inch(4.93)
style_body(body,13.1)
for p in body.text_frame.paragraphs:
    txt=p.text.strip()
    if txt in ['แนวทางป้องกันในระยะยาว']:
        for r in p.runs: r.font.size=Pt(17); r.font.bold=True; r.font.color.rgb=BLUE
    if txt.startswith('การแก้ปัญหาในระยะยาว'):
        for r in p.runs: r.font.bold=True; r.font.color.rgb=NAVY

s=prs.slides[8]; o=ORIG[8]
style_title(o[0],22)
add_card(s,0.78,1.52,11.78,5.20,LIGHT3,BORDER)
body=o[1]; body.left=inch(0.98); body.top=inch(1.72); body.width=inch(11.35); body.height=inch(4.85); style_body(body,13.1)
for p in body.text_frame.paragraphs:
    txt=p.text.strip()
    if txt.startswith('เรื่อง ') or txt.startswith('1.') or txt.startswith('2.') or txt.startswith('หมายเหตุ'):
        for r in p.runs: r.font.bold=True
    if txt.startswith('เรื่อง '):
        for r in p.runs: r.font.color.rgb=BLUE; r.font.size=Pt(15)

s=prs.slides[9]; o=ORIG[9]
style_title(o[1],23)
resize_table(o[0],0.58,1.38,12.16,1.62); style_table2(o[0],7.5,1,True,[3,4],[2])
resize_table(o[2],0.58,3.18,12.16,3.88); style_table2(o[2],6.8,1,True,[5,6],[3,4])

s=prs.slides[10]; o=ORIG[10]
style_title(o[1],21.5)
resize_table(o[0],0.58,1.38,12.16,1.55); style_table2(o[0],7.5,1,True,[3,4],[2])
sub=o[3]; sub.left=inch(0.64); sub.top=inch(3.08); sub.width=inch(8.0); sub.height=inch(0.36); style_body(sub,12.4,BLUE)
for p in sub.text_frame.paragraphs:
    for r in p.runs: r.font.bold=True
resize_table(o[2],0.58,3.48,12.16,3.56); style_table2(o[2],6.7,1,True,[5,6],[3,4])

s=prs.slides[11]; o=ORIG[11]
style_title(o[1],23)
resize_table(o[4],0.65,1.55,12.05,5.30); style_table2(o[4],8.2,1,True,[5,6],[3,4])

s=prs.slides[12]; o=ORIG[12]
style_title(o[0],23)
resize_table(o[3],0.58,1.38,12.16,1.75); style_table2(o[3],7.4,1,True,[3,4],[2])
sub=o[2]; sub.left=inch(0.64); sub.top=inch(3.26); sub.width=inch(6.5); sub.height=inch(0.34); style_body(sub,12.4,BLUE)
for p in sub.text_frame.paragraphs:
    for r in p.runs: r.font.bold=True
resize_table(o[1],0.58,3.64,12.16,3.37); style_table2(o[1],7.4,1,True,[5,6],[3,4])

s=prs.slides[13]; o=ORIG[13]
style_title(o[1],21.5)
resize_table(o[4],0.46,1.48,12.42,5.43); style_table2(o[4],7.7,1,True,[5,6],[4])
for r in range(1,len(o[4].table.rows)-1):
    o[4].table.cell(r,6).fill.solid(); o[4].table.cell(r,6).fill.fore_color.rgb=PALE_CYAN

s=prs.slides[14]; o=ORIG[14]
style_title(o[0],21.5)
add_card(s,0.58,1.50,4.00,5.43,LIGHT3,BORDER)
body=o[2]; body.left=inch(0.77); body.top=inch(1.72); body.width=inch(3.62); body.height=inch(4.96); style_body(body,12.2)
resize_table(o[1],4.78,1.50,8.04,5.43); style_table2(o[1],7.1,1)

s=prs.slides[15]; o=ORIG[15]
style_title(o[0],23)
add_card(s,0.58,1.50,5.95,5.40,LIGHT3,BORDER)
add_card(s,6.78,1.50,5.95,5.40,PALE_BLUE,BORDER)
left=o[1]; left.left=inch(0.78); left.top=inch(1.72); left.width=inch(5.55); left.height=inch(4.95); style_body(left,11.5)
right=o[2]; right.left=inch(6.98); right.top=inch(1.72); right.width=inch(5.55); right.height=inch(4.95); style_body(right,11.5)
for shape in [left,right]:
    for p in shape.text_frame.paragraphs:
        txt=p.text.strip()
        if txt in ['ปัญหาระบบบริหารงานทรัพย์สิน D-Fix เดิม','ประเด็นอื่นๆ','แนวทางที่ได้ดำเนินการแล้ว']:
            for r in p.runs: r.font.bold=True; r.font.color.rgb=BLUE; r.font.size=Pt(13.5)

s=prs.slides[16]; o=ORIG[16]
style_title(o[6],25)
for idx,y in [(4,1.34),(5,4.14)]:
    sh=o[idx]; sh.left=inch(0.62); sh.top=inch(y); sh.width=inch(1.0); sh.height=inch(0.34); style_body(sh,13.0,BLUE)
    for p in sh.text_frame.paragraphs:
        for r in p.runs: r.font.bold=True
o[2].left=inch(0.64); o[2].top=inch(1.72); o[2].width=inch(12.02); o[2].height=inch(2.14)
add_card(s,0.56,1.64,12.18,2.30,WHITE,BORDER)
set_pic(o[3],0.72,4.52,11.80,2.35)
add_card(s,0.56,4.44,12.18,2.51,WHITE,BORDER)

s=prs.slides[17]; o=ORIG[17]
style_title(o[2],22.5)
add_card(s,0.48,1.50,5.28,5.50,WHITE,BORDER)
add_card(s,5.92,1.50,6.93,5.50,WHITE,BORDER)
o[3].left=inch(0.66); o[3].top=inch(1.70); o[3].width=inch(4.92); o[3].height=inch(5.10)
set_pic(o[4],6.10,1.70,6.57,5.10)

s=prs.slides[18]; o=ORIG[18]
add_card(s,1.37,0.98,10.60,5.98,WHITE,BORDER)
o[2].left=inch(1.55); o[2].top=inch(1.14); o[2].width=inch(10.25); o[2].height=inch(5.66)

for slide, orig in zip(prs.slides, ORIG):
    move_original_content_to_front(slide, orig)

for slide in prs.slides:
    for sh in slide.shapes:
        if getattr(sh,'has_text_frame',False) and sh.shape_type in [MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER]:
            try:
                sh.fill.background(); sh.line.fill.background()
            except: pass

AFTER = snapshot_original(prs)
diffs=[]
for k,v in BEFORE.items():
    if AFTER.get(k) != v:
        diffs.append((k,v,AFTER.get(k)))
if diffs:
    raise RuntimeError(f"Text content changed. diffs={diffs[:10]}")

prs.save(OUT)
print(OUT)
print('size', os.path.getsize(OUT))
print('text entries preserved', len(BEFORE))
